#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_trace_events.py — trace 事件「声明 vs 产出」回归（v3.17.0 B3-2）

立项背景（不是"补几个事件"的例行动作）：
  config/harness.yaml#observability.trace_events 声明了 12 个 trace 事件，
  但 v3.17.0 独立复核实测**只有 5 个有代码产出**（gate / stage_enter / stage_exit /
  budget_exceeded / no_progress），另外 7 个是**纯声明**：
      team_create / preflight / spawn / spawn_failed / degrade / member_return / error

  代价是立项级的：同一份 yaml 的 fallback.require_disclosure=true 写着
  "降级必须在交付物与 trace 中显式声明"，而 degrade 事件**没有任何代码产出**
  ⇒ "降级是否发生、是否被声明"在 trace 里根本查不到。
  这正是 v2.2.0 建立整个 harness 层的原始理由（"静默降级无人察觉"）被自己重犯。

本文件把这件事钉成机器断言：
  ① 12 个声明的事件**逐个**有真实产出（跑真命令 → 读真 trace 文件 → 断言事件在）；
  ② trace report 能报出"声明了但没产出"的差集（declared_events_missing）；
  ③ trace report 能报出"产出了但没声明"的差集（undeclared_events）——
     这是同型反向缺口（agent_spawn 曾是这种：代码在产、yaml 没登记）；
  ④ degrade 事件必须带齐"成员 id / 降级通道 / 原因"三要素——
     这是 fallback.require_disclosure 的机器可查形态。

用法：
    python3 -m pytest scripts/test_trace_events.py -q     # pytest 收集
    python3 scripts/test_trace_events.py                  # 直跑（ensure_settings 第 7 项会跑）
退出码（直跑）：0=全通过 1=有用例失败
"""

import json
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
HARNESS = os.path.join(HERE, "harness.py")

PASS, FAIL = [], []

# yaml#observability.trace_events 声明的全部 12 个事件（唯一源 = config/harness.yaml）
DECLARED_EVENTS = ("team_create", "preflight", "spawn", "spawn_failed", "degrade",
                   "member_return", "gate", "stage_enter", "stage_exit",
                   "budget_exceeded", "error", "no_progress")


def run(args, cwd=None):
    p = subprocess.run([sys.executable, HARNESS] + args, capture_output=True, text=True, cwd=cwd)
    return p.returncode, p.stdout, p.stderr


def check(name, cond, extra=""):
    """唯一断言口：记 PASS/FAIL，失败时抛 AssertionError（pytest 可见）。"""
    if cond:
        PASS.append(name)
        print(f"  ✅ {name}")
    else:
        FAIL.append(name)
        msg = f"{name}  {extra}"
        print(f"  ❌ {msg}")
        raise AssertionError(msg)


def _trace_events(ws):
    """读工作区里的 trace 文件 → 事件名列表（保序）。"""
    p = os.path.join(ws, ".workbuddy", "harness_trace.jsonl")
    if not os.path.exists(p):
        return []
    out = []
    for line in open(p, encoding="utf-8"):
        line = line.strip()
        if not line:
            continue
        try:
            out.append(json.loads(line).get("event"))
        except Exception:
            pass
    return out


def _trace_records(ws):
    """读工作区里的 trace 文件 → 完整记录列表。"""
    p = os.path.join(ws, ".workbuddy", "harness_trace.jsonl")
    if not os.path.exists(p):
        return []
    out = []
    for line in open(p, encoding="utf-8"):
        line = line.strip()
        if not line:
            continue
        try:
            out.append(json.loads(line))
        except Exception:
            pass
    return out


def _ws():
    return tempfile.TemporaryDirectory()


# ══════════════════════════ 用例 ══════════════════════════
def _make_deck(path):
    """造一份最小 pptx（gate 事件需要真课件）；python-pptx 不可用时退化为占位文件。"""
    try:
        from pptx import Presentation
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(path)
        return True
    except Exception:
        with open(path, "wb") as f:
            f.write(b"PK\x03\x04fake-deck-for-trace-events-test")
        return False


def test_声明的12个事件全部有代码产出():
    """逐个跑真命令 → 读真 trace 文件 → 断言该事件在文件里。

    这是本任务的核心断言：**声明了却没有产出点**的事件，任何一个都会让本用例红。
    """
    with _ws() as w:
        # team_create：路由判出要建队（L1+）→ 写 team_create
        run(["route", "--level", "L2", "--workspace", w])
        # preflight：开工前可用性预检 → 写 preflight
        run(["preflight", "--team", "vet-med-ppt-team", "--task-id", "T1", "--workspace", w])
        # spawn：spawn 成功 → 写 spawn
        run(["spawn", "--agent", "dvm-content-expert", "--task-id", "T1", "--stage", "1",
             "--input-bytes", "1200", "--output-bytes", "800", "--workspace", w])
        # spawn_failed：spawn 失败 → 写 spawn_failed
        run(["spawn-failed", "--agent", "pedagogy-architect", "--task-id", "T1", "--stage", "1",
             "--attempt", "2", "--reason", "平台返回 not available", "--workspace", w])
        # degrade：成员降级 → 写 degrade
        run(["degrade", "--agent", "pedagogy-architect", "--task-id", "T1", "--stage", "1",
             "--channel", "general-purpose+persona", "--reason", "spawn 2 次均不可用", "--workspace", w])
        # member_return：成员回传 → 写 member_return
        run(["member-return", "--agent", "dvm-content-expert", "--task-id", "T1", "--stage", "1",
             "--status", "success", "--artifact", "out/blueprint.md", "--workspace", w])
        # stage_enter / stage_exit：既有产出点
        run(["stage", "--action", "enter", "--stage", "1", "--task-id", "T1", "--workspace", w])
        run(["stage", "--action", "exit", "--stage", "1", "--task-id", "T1", "--workspace", w])
        # budget_exceeded：超限即写
        run(["budget", "reset", "--kind", "spawn", "--workspace", w])
        run(["budget", "spend", "--kind", "spawn", "--n", "99", "--workspace", w])
        # no_progress：连续 N 轮 P0+P1 未下降 → 写（loop continue 路径）
        deck = os.path.join(w, "demo.pptx")
        _make_deck(deck)
        run(["loop", "start", "--deck", deck, "--workspace", w])
        for p0 in (3, 3, 3):
            run(["loop", "continue", "--decision", "accept", "--confirmed",
                 "--p0", str(p0), "--p1", "2", "--p2", "5", "--allow-unverified",
                 "--workspace", w])
        # gate：既有产出点（需真课件；用最小 pptx 触发 verify_deck → 写 gate 事件）
        gate_deck = os.path.join(w, "gate_demo.pptx")
        _make_deck(gate_deck)
        run(["gate", "--deck", gate_deck, "--hours", "1", "--level", "P0", "--no-snap",
             "--workspace", w])
        # error：顶层异常出口（经 run_command 注入异常——异常路径不能只藏在 __main__ 里）
        sys.path.insert(0, HERE)
        import harness as _h

        class _Boom:
            workspace = w
            func = staticmethod(lambda a: (_ for _ in ()).throw(RuntimeError("注入的运行期异常")))
        _h.run_command(_Boom())

        evs = _trace_events(w)
        missing = [e for e in DECLARED_EVENTS if e not in evs]
        check(f"trace：12 个声明事件全部有产出（实际产出 {sorted(set(evs))}）",
              not missing, f"缺产出：{missing}")


def test_声明与产出双向差集可查():
    """trace report 必须报出两个方向的差集。

    ① declared_events_missing —— "声明了却没产出"（本轮治的病）
    ② undeclared_events     —— "产出了却没声明"（反向同型病；agent_spawn 曾是这种）
    """
    with _ws() as w:
        run(["spawn", "--agent", "dvm-content-expert", "--task-id", "T1", "--workspace", w])
        rc, out, _ = run(["trace", "report", "--workspace", w])
        try:
            rep = json.loads(out)
        except Exception:
            check("trace report 可解析", False, out[:200])
            return
        check("trace report 输出 declared_events（yaml 声明的全集）",
              set(rep.get("declared_events") or []) == set(DECLARED_EVENTS),
              f"got {rep.get('declared_events')}")
        check("trace report 输出 declared_events_missing（声明了却没产出的）",
              isinstance(rep.get("declared_events_missing"), list),
              f"got {rep.get('declared_events_missing')!r}")
        check("trace report 输出 undeclared_events（产出了却没声明的）",
              isinstance(rep.get("undeclared_events"), list),
              f"got {rep.get('undeclared_events')!r}")
        # 本工作区只跑了 spawn ⇒ 其余 11 个声明事件应被报为 missing（差集确实在算）
        miss = rep.get("declared_events_missing") or []
        check("trace report 的 missing 差集确实反映未产出事件",
              "gate" in miss and "spawn" not in miss, f"got {miss}")


def test_degrade事件带齐成员id通道原因三要素():
    """fallback.require_disclosure 的机器可查形态：三要素缺一不可。

    缺了任何一个，下游就无法判断"降级了什么、降成了什么、为什么"，
    require_disclosure 又退回"只有人读"。
    """
    with _ws() as w:
        run(["degrade", "--agent", "pedagogy-architect", "--task-id", "T9", "--stage", "2",
             "--channel", "general-purpose+persona", "--reason", "spawn 2 次均 not available",
             "--workspace", w])
        recs = [r for r in _trace_records(w) if r.get("event") == "degrade"]
        check("degrade 事件已写出", len(recs) == 1, f"got {len(recs)}")
        d = recs[0].get("detail") or {}
        check("degrade 事件含成员 id", d.get("member_id") == "pedagogy-architect",
              f"got {d.get('member_id')!r}")
        check("degrade 事件含降级通道", d.get("degrade_channel") == "general-purpose+persona",
              f"got {d.get('degrade_channel')!r}")
        check("degrade 事件含原因", bool(d.get("reason")), f"got {d.get('reason')!r}")
        check("degrade 事件标明 require_disclosure 为真（声明被执行）",
              d.get("disclosure_required") is True, f"got {d.get('disclosure_required')!r}")


def test_member_return事件含成员task_id状态与产物路径():
    """成员回传要可审计：谁 / 哪个任务 / 什么状态 / 产物在哪。"""
    with _ws() as w:
        run(["member-return", "--agent", "dvm-content-expert", "--task-id", "T7", "--stage", "1",
             "--status", "partial", "--artifact", "out/evidence.md",
             "--issues", "缺 2 条证据", "--workspace", w])
        recs = [r for r in _trace_records(w) if r.get("event") == "member_return"]
        check("member_return 事件已写出", len(recs) == 1, f"got {len(recs)}")
        d = recs[0].get("detail") or {}
        check("member_return 含成员 id（actor）",
              recs[0].get("actor") == "dvm-content-expert", f"got {recs[0].get('actor')!r}")
        check("member_return 含 task_id", d.get("task_id") == "T7", f"got {d.get('task_id')!r}")
        check("member_return 含 status", d.get("status") == "partial", f"got {d.get('status')!r}")
        check("member_return 含产物路径", d.get("artifact") == "out/evidence.md",
              f"got {d.get('artifact')!r}")
        check("member_return 含 issues", d.get("issues") == "缺 2 条证据",
              f"got {d.get('issues')!r}")


def test_spawn事件含成员id与任务id():
    """spawn 成功事件要写明成员 id / 任务 id（yaml 声明的事件名，不是 agent_spawn）。"""
    with _ws() as w:
        run(["spawn", "--agent", "pedagogy-architect", "--task-id", "T3", "--stage", "2",
             "--input-bytes", "1000", "--output-bytes", "500", "--workspace", w])
        evs = _trace_events(w)
        check("spawn 事件已写出（用 yaml 声明的事件名）", "spawn" in evs, f"got {evs}")
        recs = [r for r in _trace_records(w) if r.get("event") == "spawn"]
        d = (recs[0].get("detail") or {}) if recs else {}
        check("spawn 事件含成员 id", recs and recs[0].get("actor") == "pedagogy-architect",
              f"got {recs[0].get('actor') if recs else None!r}")
        check("spawn 事件含任务 id", d.get("task_id") == "T3", f"got {d.get('task_id')!r}")
        check("spawn 保留 agent_spawn 兼容事件（既有文档引用它，删了即成假引用）",
              "agent_spawn" in evs, f"got {evs}")


def test_spawn_failed事件含尝试次数与上限():
    """spawn 失败要写明试了几次 / 上限是多少 / 为什么不行。"""
    with _ws() as w:
        run(["spawn-failed", "--agent", "visual-slide-designer", "--task-id", "T4", "--stage", "2",
             "--attempt", "2", "--reason", "not available", "--workspace", w])
        recs = [r for r in _trace_records(w) if r.get("event") == "spawn_failed"]
        check("spawn_failed 事件已写出", len(recs) == 1, f"got {len(recs)}")
        d = (recs[0].get("detail") or {}) if recs else {}
        check("spawn_failed 含尝试次数", d.get("attempt") == 2, f"got {d.get('attempt')!r}")
        check("spawn_failed 含尝试上限（取自 fallback.max_spawn_attempts）",
              d.get("max_attempts") == 2, f"got {d.get('max_attempts')!r}")
        check("spawn_failed 标明是否已耗尽", d.get("exhausted") is True,
              f"got {d.get('exhausted')!r}")
        check("spawn_failed 含失败原因", bool(d.get("reason")), f"got {d.get('reason')!r}")


def test_preflight检出静默降级时写degrade事件():
    """preflight 是机器唯一能抓到"平台把成员回落成通用子代理"的时点。

    抓到就必须留痕（否则"静默降级"仍然是静默的）。
    """
    with _ws() as w:
        # 造一个"成员被平台登记为 general-purpose"的团队配置
        teams = os.path.expanduser("~/.workbuddy/teams/vet-med-ppt-team")
        os.makedirs(teams, exist_ok=True)
        cfgp = os.path.join(teams, "config.json")
        backup = None
        if os.path.exists(cfgp):
            backup = open(cfgp, encoding="utf-8").read()
        try:
            with open(cfgp, "w", encoding="utf-8") as f:
                json.dump({"members": [
                    {"name": "dvm-content-expert", "agentType": "general-purpose"},
                    {"name": "pedagogy-architect", "agentType": "pedagogy-architect"},
                ]}, f)
            rc, _, _ = run(["preflight", "--team", "vet-med-ppt-team", "--task-id", "T5",
                            "--workspace", w])
            recs = [r for r in _trace_records(w) if r.get("event") == "degrade"]
            check("preflight 检出静默降级 → 写 degrade 事件", len(recs) >= 1, f"got {len(recs)}")
            d = (recs[0].get("detail") or {}) if recs else {}
            check("preflight 的 degrade 事件标明检出来源是 preflight（不是人手工申报）",
                  d.get("detected_by") == "preflight", f"got {d.get('detected_by')!r}")
            check("preflight 的 degrade 事件含降级通道（取自 members[*].fallback）",
                  bool(d.get("degrade_channel")), f"got {d.get('degrade_channel')!r}")
            check("preflight 事件本身也写出（yaml 声明的事件）",
                  "preflight" in _trace_events(w), f"got {_trace_events(w)}")
        finally:
            if backup is not None:
                with open(cfgp, "w", encoding="utf-8") as f:
                    f.write(backup)
            else:
                try:
                    os.remove(cfgp)
                except OSError:
                    pass


def test_error事件在运行期异常时写出():
    """运行期异常此前只打栈、trace 里查不到"哪一步炸了" ⇒ 编排现场无法复盘。

    触发方式：直接调 `harness.run_command()` 并塞一个必然抛异常的子命令——
    异常路径若只藏在 __main__ 里就永远测不到（这正是本包反复吃亏的"不可达路径"型缺陷）。
    """
    sys.path.insert(0, HERE)
    import harness as _h
    with _ws() as w:
        class _Args:
            workspace = w
            func = staticmethod(lambda a: (_ for _ in ()).throw(RuntimeError("注入的运行期异常")))
        rc = _h.run_command(_Args())
        recs = [r for r in _trace_records(w) if r.get("event") == "error"]
        check("运行期异常 → 写出 error 事件", len(recs) == 1, f"got {len(recs)}")
        d = (recs[0].get("detail") or {}) if recs else {}
        check("error 事件含子命令名", bool(d.get("command")), f"got {d.get('command')!r}")
        check("error 事件含异常类型", d.get("exception") == "RuntimeError",
              f"got {d.get('exception')!r}")
        check("error 事件含异常消息", "注入的运行期异常" in (d.get("message") or ""),
              f"got {d.get('message')!r}")
        check("error 事件含 argv（可复现）", isinstance(d.get("argv"), list),
              f"got {d.get('argv')!r}")
        check("运行期异常：退出码为 4（运行错误，非课件质量）", rc == 4, f"got rc={rc}")


def test_run_command正常路径不写error事件():
    """反向断言：没异常就不写 error 事件（否则 error 事件会变成噪声）。"""
    sys.path.insert(0, HERE)
    import harness as _h
    with _ws() as w:
        class _Args:
            workspace = w
            func = staticmethod(lambda a: 0)
        rc = _h.run_command(_Args())
        check("正常路径：退出码透传", rc == 0, f"got rc={rc}")
        check("正常路径：不写 error 事件", "error" not in _trace_events(w),
              f"got {_trace_events(w)}")


def test_L0路由不写team_create事件():
    """L0 不建队 → 不得写 team_create。

    写了就是"声明了却没发生"的假事件——正是本任务要治的病。
    """
    with _ws() as w:
        run(["route", "--level", "L0", "--workspace", w])
        evs = _trace_events(w)
        check("L0（不建队）不写 team_create 事件", "team_create" not in evs, f"got {evs}")


def main():
    ap = __import__("argparse").ArgumentParser()
    ap.add_argument("--deck", default="", help="（本文件不消费，占位以兼容统一 runner）")
    a = ap.parse_args()

    print("== 回归测试：trace 事件「声明 vs 产出」v3.17.0 B3-2 ==")
    for name, fn in list(globals().items()):
        if not (name.startswith("test_") and callable(fn)):
            continue
        try:
            fn()
        except AssertionError:
            pass          # 已由 check() 记入 FAIL
        except Exception as e:
            FAIL.append(name)
            print(f"  ❌ {name} 用例异常：{type(e).__name__}: {e}")

    print(f"\n通过 {len(PASS)} / 失败 {len(FAIL)}")
    if FAIL:
        print("失败用例：", FAIL)
        return 1
    print("✅ 全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
