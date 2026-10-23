#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_harness.py — 编排内核与门禁的回归测试（v2.2.0 新增）

背景：本包此前 0 测试、0 CI。2026-09-29 基线实测在一个真实课件上撞出 3 个
未被发现的崩溃（越界/重叠/字号比较遇 None），全部是"没有回归测试"的直接后果。
本文件提供冒烟级回归：不追求覆盖率，只锁死"曾经出过事的那几条路径"。

v3.16.0（假绿修复 —— 实测事故，不是理论风险）：
  此前全部断言写在 `if __name__ == "__main__"` 里，**没有**任何 `def test_*`，
  于是 `python3 -m pytest scripts/ -q` 收集到 0 条用例、照样显示全绿；
  只有 ensure_settings 第 7 项（直跑套件）才把它们拉出来。
  后果（本轮真实发生两次）：F3 内容判据一上线，不合格夹具被判失败，
  pytest 侧**毫无反应**，两名成员都误以为回归全过。
  修法：每个断言改成独立的 `def test_*()` —— pytest 与直跑共用同一批用例；
  `check()` 失败即抛 AssertionError（pytest 可见），直跑时由 main() 逐个捕获汇总。

用法：
    python3 -m pytest scripts/test_harness.py -q        # pytest 收集（推荐）
    python3 scripts/test_harness.py                     # 直跑全量（ensure_settings 第 7 项用）
    python3 scripts/test_harness.py --deck <课件.pptx>  # 额外跑真实课件门禁冒烟（可选）
    HARNESS_TEST_DECK=<课件.pptx> python3 -m pytest scripts/test_harness.py -q
退出码（直跑）：0=全通过 1=有用例失败
"""

import json
import os
import subprocess
import sys
import tempfile
import time

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
HARNESS = os.path.join(HERE, "harness.py")
VERIFY = os.path.join(HERE, "verify_deck.py")

PASS, FAIL = [], []
# 真实课件门禁冒烟：直跑用 --deck，pytest 用环境变量（缺省即跳过，不算失败）
DECK = os.environ.get("HARNESS_TEST_DECK", "")


def run(args, cwd=None):
    p = subprocess.run([sys.executable, HARNESS] + args, capture_output=True, text=True, cwd=cwd)
    return p.returncode, p.stdout, p.stderr


def check(name, cond, extra=""):
    """唯一断言口：记 PASS/FAIL，并在失败时抛 AssertionError。

    为什么要抛：只 print 不抛的话 pytest 会把它算成"通过"（本文件假绿的成因）。
    """
    if cond:
        PASS.append(name)
        print(f"  ✅ {name}")
    else:
        FAIL.append(name)
        msg = f"{name}  {extra}"
        print(f"  ❌ {msg}")
        raise AssertionError(msg)


def _ws():
    """每个用例独立工作区（pytest 与直跑共用；避免用例之间互相污染状态文件）。"""
    return tempfile.TemporaryDirectory()


def _env_fixture(w, name, obj):
    """在工作区里落一个含 result_envelope JSON 块的 .md 文件。"""
    p = os.path.join(w, f"env_{name}.md")
    with open(p, "w", encoding="utf-8") as f:
        f.write("```json\n" + json.dumps(obj, ensure_ascii=False) + "\n```\n")
    return p


# 合规基线信封：confidence 带文件路径+计数，deliverable_summary 带落盘路径+原始 P0/P1/P2 计数
OK_ENV = {"agent_id": "verifier", "task_id": "T1", "phase": "3", "status": "success",
          "confidence": "0.85（依据：scripts/verify_deck.py 复算 P0=0/P1=2，共 6 页）",
          "deliverable_summary": "重写 6 页正文，落盘 demo_wip/6 页；P0=0 / P1=2 / P2=5"}


# ══════════════════════════ 用例 ══════════════════════════
def test_route_L2_输出4成员且带预算():
    rc, out, _ = run(["route", "--level", "L2"])
    try:
        plan = json.loads(out)
        ok = rc == 0 and len(plan.get("spawn", [])) == 4 and "budgets" in plan
    except Exception:
        ok = False
    check("route L2 输出 4 成员且带预算", ok, out[:200])


def test_state_init_validate_通过():
    with _ws() as w:
        rc, _, _ = run(["state", "init", "--workspace", w])
        rc2, out2, _ = run(["state", "validate", "--workspace", w])
        check("state init + validate 通过", rc == 0 and rc2 == 0, out2[:200])


def test_state_非法阶段9被拒绝():
    # 曾出现 current_stage="M5A2" 这类便签值
    with _ws() as w:
        rc, _, err = run(["state", "set", "--stage", "9", "--workspace", w])
        check("非法阶段 9 被拒绝", rc == 1, err[:120])


def test_state_白名单外键被拒绝():
    # 曾膨胀到 32 键 / 27KB
    with _ws() as w:
        rc, _, err = run(["state", "set", "--key", "queue_case_17", "--value", "[]",
                          "--workspace", w])
        check("白名单外键被拒绝", rc == 1, err[:120])


def test_budget_第8次放行第9次阻断():
    with _ws() as w:
        run(["budget", "reset", "--kind", "spawn", "--workspace", w])
        last_ok, last_block = None, None
        for i in range(9):
            rc, _, _ = run(["budget", "spend", "--kind", "spawn", "--n", "1", "--workspace", w])
            if i == 7:
                last_ok = rc
            if i == 8:
                last_block = rc
        check("预算：第 8 次放行、第 9 次阻断", last_ok == 0 and last_block == 1,
              f"got ok={last_ok} block={last_block}")


def test_trace_落盘并可按事件汇总():
    with _ws() as w:
        run(["trace", "log", "--event", "spawn", "--actor", "pedagogy-architect",
             "--detail", '{"stage":3}', "--workspace", w])
        run(["trace", "log", "--event", "spawn_failed", "--actor", "pedagogy-architect",
             "--detail", '{"reason":"not available"}', "--workspace", w])
        rc, out, _ = run(["trace", "report", "--workspace", w])
        try:
            rep = json.loads(out)
            ok = rc == 0 and rep.get("by_event", {}).get("spawn") == 1
        except Exception:
            ok = False
        check("trace 落盘并可按事件汇总", ok, out[:200])


def test_stage_enter_exit_配对写入stage_seconds():
    with _ws() as w:
        rc_e, _, _ = run(["stage", "--action", "enter", "--stage", "1",
                          "--task-id", "T-stage", "--workspace", w])
        time.sleep(1.1)
        rc_x, _, _ = run(["stage", "--action", "exit", "--stage", "1",
                          "--task-id", "T-stage", "--workspace", w])
        rc_r, out_r, _ = run(["trace", "report", "--workspace", w])
        try:
            rep2 = json.loads(out_r)
            secs = rep2.get("stage_seconds", {}).get("1", [])
            ok = (rc_e == 0 and rc_x == 0 and rc_r == 0
                  and len(secs) == 1 and secs[0] >= 1
                  and rep2.get("stage_seconds_status") == "ok")
        except Exception:
            ok = False
        check("stage enter/exit 配对写入 stage_seconds", ok, out_r[:200])


def test_stage_边界_非法stage拒绝与孤儿exit警告():
    with _ws() as w:
        rc_bad, _, _ = run(["stage", "--action", "enter", "--stage", "9", "--workspace", w])
        rc_orphan, _, err_orphan = run(["stage", "--action", "exit", "--stage", "3",
                                        "--workspace", w])
        check("stage 边界：stage=9 拒绝（exit 2）、孤儿 exit 有警告",
              rc_bad == 2 and rc_orphan == 0 and "孤儿" in err_orphan,
              f"got bad={rc_bad} orphan={rc_orphan} err={err_orphan[:80]}")


def test_preflight_识别全部4成员MD():
    with _ws() as w:
        rc, out, _ = run(["preflight", "--workspace", w])
        try:
            pf = json.loads(out)
            ok = rc == 0 and all(v.get("md_exists") for v in pf.get("members", {}).values())
        except Exception:
            ok = False
        check("preflight 识别全部 4 成员 MD", ok, out[:200])


def test_verify_artifact_不存在拒_存在过_体积不符告警():
    # 实测 3 名成员自报"已落盘"而文件不存在
    with _ws() as w:
        rc, _, _ = run(["verify-artifact", "--path", os.path.join(w, "不存在的产物.md"),
                        "--workspace", w])
        real = os.path.join(w, "real.md")
        with open(real, "w", encoding="utf-8") as f:
            f.write("x" * 500)
        rc2, _, _ = run(["verify-artifact", "--path", real, "--expect-bytes", "500",
                         "--workspace", w])
        rc3, _, _ = run(["verify-artifact", "--path", real, "--expect-bytes", "99999",
                         "--workspace", w])
        check("产物核验：不存在→拒、存在→过、体积对不上→告警",
              rc == 1 and rc2 == 0 and rc3 == 1, f"got {rc},{rc2},{rc3}")


def test_envelope_合规通过_缺字段退回():
    # v3.16.0：原 good 夹具是 "confidence":80 / "deliverable_summary":"ok"——
    #   在只判"字段存在"的年代能过，但它本身就是 yaml#fields 明令禁止的裸标签/裸表述
    #   （fields.confidence.rule 禁裸标签、fields.deliverable_summary.rule 要求可验证计数）。
    #   内容判据上线后**不能**反过来放宽规则去迁就夹具——那样等于把新判据废掉；
    #   正确修法是把夹具改成 yaml 要求的合规写法（带依据 + 落盘路径 + 原始计数）。
    with _ws() as w:
        good = _env_fixture(w, "good", OK_ENV)
        bad = _env_fixture(w, "bad", {"agent_id": "x", "status": "成功"})
        rc_e1, _, _ = run(["verify-envelope", "--path", good, "--workspace", w])
        rc_e2, _, _ = run(["verify-envelope", "--path", bad, "--workspace", w])
        check("信封校验：合规→过、缺字段→退回", rc_e1 == 0 and rc_e2 == 1,
              f"got {rc_e1},{rc_e2}")


def test_content_rule_confidence_裸标签必失败():
    with _ws() as w:
        p = _env_fixture(w, "bare_conf", dict(OK_ENV, confidence="高"))
        rc, so, _ = run(["verify-envelope", "--path", p, "--workspace", w])
        check("内容判据：confidence='高' 裸标签 → 判失败（exit 1）", rc == 1, f"got {rc} {so[:160]}")
        check("内容判据：报错引用 yaml 规则路径（fields.confidence.rule + content_rules）",
              "fields.confidence.rule" in so
              and "result-envelope.yaml#validation.content_rules" in so, f"got {so[:200]}")


def test_content_rule_confidence_带依据必通过():
    # 三种合规依据形态：文件路径+计数 / 命令输出（反引号）/ 工具名+中文计数
    variants = {
        "路径+计数": "0.85（依据：scripts/verify_deck.py 复算 P0=0/P1=2，共 6 页）",
        "命令输出": "依据 `python3 scripts/verify_deck.py --hours 2` 输出判定，P0=0",
        "工具名+计数": "依据本次 lint 复算，缺陷由 3 处降至 0 处",
    }
    with _ws() as w:
        for name, conf in variants.items():
            p = _env_fixture(w, f"ok_conf_{abs(hash(name))}", dict(OK_ENV, confidence=conf))
            rc, so, _ = run(["verify-envelope", "--path", p, "--workspace", w])
            check(f"内容判据：confidence 带依据（{name}）→ 通过（exit 0）", rc == 0,
                  f"got {rc} {so[:200]}")


def test_content_rule_deliverable_summary_裸表述必失败():
    bare_cases = {"已优化": "已优化", "全部完成": "全部完成", "归零类": "问题已修复归零"}
    with _ws() as w:
        for name, summ in bare_cases.items():
            p = _env_fixture(w, f"bare_sum_{abs(hash(name))}",
                             dict(OK_ENV, deliverable_summary=summ))
            rc, so, _ = run(["verify-envelope", "--path", p, "--workspace", w])
            check(f"内容判据：deliverable_summary='{summ}'（{name}，无路径无计数）→ 判失败",
                  rc == 1, f"got {rc} {so[:200]}")
            check(f"内容判据：{name} 报错引用 fields.deliverable_summary.rule",
                  "fields.deliverable_summary.rule" in so, f"got {so[:200]}")


def test_content_rule_deliverable_summary_含路径与计数必通过():
    ok_cases = {
        "路径+计数": "重写 6 页正文，落盘 demo_wip/v3.pptx；P0=0 / P1=2 / P2=5",
        "中文计数": "已交付三页新版正文，落盘 demo_wip/第3页.md",
    }
    with _ws() as w:
        for name, summ in ok_cases.items():
            p = _env_fixture(w, f"ok_sum_{abs(hash(name))}",
                             dict(OK_ENV, deliverable_summary=summ))
            rc, so, _ = run(["verify-envelope", "--path", p, "--workspace", w])
            check(f"内容判据：deliverable_summary（{name}）→ 通过（exit 0）", rc == 0,
                  f"got {rc} {so[:200]}")


def test_content_rule_status_success_issues非空判失败():
    # 取舍记录：本条选 **失败**（而非告警）。理由——
    #   ① fields.status.criteria.success 明写"完整满足 objective，且 issues 为空"，
    #      issues 非空还自评 success 是**自评失真**，正是回传信封要杜绝的行为；
    #   ② 告警会被下游当"可放行"，失真就拦不住；
    #   ③ 判据完全确定（issues 是否非空是布尔事实），无假阳性空间。
    with _ws() as w:
        p = _env_fixture(w, "lie_status",
                         dict(OK_ENV, issues=[{"kind": "evidence_gap", "detail": "缺 2 条 PMID"}]))
        rc, so, _ = run(["verify-envelope", "--path", p, "--workspace", w])
        check("内容判据：status=success 且 issues 非空 → 判失败（取舍：失败而非告警）",
              rc == 1, f"got {rc} {so[:200]}")
        check("内容判据：status 一致性报错引用 fields.status.criteria",
              "fields.status.criteria" in so, f"got {so[:200]}")


def test_content_rule_metrics_未知键只告警不失败():
    # 取舍记录：本条选 **告警**（severity=warn）。理由——metrics 是选填自增字段，
    #   未登记指标属"可追溯性瑕疵"，不造成自评失真；判失败会逼成员删指标或谎报键名
    #   （比不登记更糟），故只留痕由主理人决定是否补进 fields.metrics.keys。
    with _ws() as w:
        p = _env_fixture(w, "metrics", dict(OK_ENV, metrics={"duration_ms": 1200,
                                                             "retry_count": 1, "new_metric": 7}))
        rc, so, se = run(["verify-envelope", "--path", p, "--workspace", w])
        check("内容判据：metrics 未知键 → 告警且不判失败（exit 0 + stderr 留痕）",
              rc == 0 and "new_metric" in (so + se), f"got {rc} {se[:160]}")


def test_three_source_阈值唯一源为domain_config():
    # v3.16.0 F2：harness.yaml#budgets 已删 three_source_* 两键，改读 domain_config.yaml
    with _ws() as w:
        rc_t, so_t, _ = run(["budget", "check", "--kind", "three_source_pass", "--workspace", w])
        try:
            import yaml as _y
            dc = _y.safe_load(open(os.path.join(PKG, "config", "domain_config.yaml"),
                                   encoding="utf-8")) or {}
            exp = (dc.get("three_source_check") or {}).get("pass_threshold")
            hy = _y.safe_load(open(os.path.join(PKG, "config", "harness.yaml"),
                                   encoding="utf-8")) or {}
            hy_bud = hy.get("budgets") or {}
            t = json.loads(so_t)
            ok = (rc_t == 0 and t.get("value") == exp
                  and "domain_config.yaml#three_source_check.pass_threshold" in t.get("config_source", "")
                  and "three_source_pass_threshold" not in hy_bud
                  and "three_source_sample_rate" not in hy_bud)
        except Exception as e:
            ok = False
            print("   （跳过三源阈值核对：", e, "）")
        check("三源阈值读 domain_config.yaml 且 harness.yaml 已无重复键", ok, f"got {so_t[:200]}")


def test_preflight_抓出静默降级为general_purpose():
    teams = os.path.expanduser("~/.workbuddy/teams")
    fake = os.path.join(teams, "_harness_selftest")
    os.makedirs(fake, exist_ok=True)
    with _ws() as w:
        try:
            with open(os.path.join(fake, "config.json"), "w", encoding="utf-8") as f:
                json.dump({"name": "_harness_selftest",
                           "members": [{"name": "pedagogy-architect",
                                        "agentType": "general-purpose"}]}, f)
            rc, out, _ = run(["preflight", "--team", "_harness_selftest", "--workspace", w])
            pf = json.loads(out)
            ok = rc == 1 and pf["members"]["pedagogy-architect"]["status"] == "degraded"
        except Exception as e:
            ok = False
            print("   （跳过降级探测：", e, "）")
        finally:
            try:
                os.remove(os.path.join(fake, "config.json"))
                os.rmdir(fake)
            except OSError:
                pass
        check("preflight 能抓出「静默降级为 general-purpose」", ok)


def test_verify_deck_机器草稿占比与judge溯源():
    # 裁决④：机器草稿计入覆盖率但单独统计——用合成样本，不依赖真实课件
    print("-- 机器草稿占比（合成样本）")
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        print("   （跳过：python-pptx 未安装）")
        return
    d = tempfile.mkdtemp()
    synth = os.path.join(d, "synth_note.pptx")
    prs = Presentation()
    for i in range(3):
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        tf = sl.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame
        tf.text = f"第{i+1}页"
        tf.paragraphs[0].runs[0].font.size = Pt(24)
        if i == 0:
            sl.notes_slide.notes_text_frame.text = "教师人工讲稿：本页定位与讲法。"
        if i == 1:
            sl.notes_slide.notes_text_frame.text = "【AI草稿·待教师复核】机器草稿…"
    prs.save(synth)
    out = os.path.join(d, "r.json")
    subprocess.run([sys.executable, VERIFY, synth, "--hours", "1",
                    "--level", "all", "--json", out], capture_output=True, text=True)
    r = json.load(open(out, encoding="utf-8"))
    g = r["gate"]
    check("机器草稿单独计数且计入覆盖率（2 页备注 / 1 页机器 = 50%）",
          g.get("machine_note_pages") == 1 and abs(g.get("machine_note_ratio", 0) - 0.5) < 1e-6,
          f"got pages={g.get('machine_note_pages')} ratio={g.get('machine_note_ratio')}")
    check("judge 溯源字段落盘（SSOT + 版本 + 课时口径）",
          r.get("judge", {}).get("role") == "SSOT" and r.get("judge", {}).get("hours_arg") == 1.0,
          json.dumps(r.get("judge"), ensure_ascii=False)[:120])
    p2 = subprocess.run([sys.executable, VERIFY, synth], capture_output=True, text=True)
    check("--hours 缺省被拒（防默认值凭空判课时不达标）", p2.returncode != 0)


def test_verify_deck_真实课件门禁冒烟_需HARNESS_TEST_DECK():
    # 退出码不得为 4（运行错误）。需真实课件：直跑 --deck <path> 或环境变量；缺省跳过。
    if not DECK:
        print("-- 门禁冒烟：未提供 --deck / HARNESS_TEST_DECK，跳过")
        return
    print(f"-- 门禁冒烟：{os.path.basename(DECK)}")
    p = subprocess.run([sys.executable, VERIFY, DECK, "--hours", "2", "--level", "all"],
                       capture_output=True, text=True)
    check(f"verify_deck 未崩溃（exit={p.returncode}，4=运行错误）", p.returncode != 4,
          (p.stderr or "")[:200])


# ══════════════════════════ 直跑入口 ══════════════════════════
def main():
    ap = __import__("argparse").ArgumentParser()
    ap.add_argument("--deck", default="", help="可选：额外跑门禁冒烟的课件路径")
    a = ap.parse_args()
    global DECK
    if a.deck:
        DECK = a.deck

    print("== 回归测试：vet-med-ppt-team harness v3.16.0（pytest 可发现）==")
    # 按定义顺序跑全部 test_*（check() 失败抛 AssertionError，这里逐个捕获继续跑）
    for name, fn in list(globals().items()):
        if not (name.startswith("test_") and callable(fn)):
            continue
        try:
            fn()
        except AssertionError:
            pass          # 已由 check() 记入 FAIL
        except Exception as e:
            FAIL.append(name)
            print(f"  ❌ {name} 用例异常：{type(e).__name__}: {e}")

    print(f"\n通过 {len(PASS)} / 失败 {len(FAIL)}")
    if FAIL:
        print("失败用例：", FAIL)
        return 1
    print("✅ 全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
