#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sample_fixtures.py — 外部课件样本的**可解析定位**层（v3.24.0）

背景（FEAT-20261023-003c 的根因取证）：
  `test_ooxml_lint.py::T1/T11` 与 `test_extract_assets.py::T6` 长期依赖**包外**课件：
    - `养禽与禽病防治/.harness-runs/M4-M7_重诊断/sample_extract/*.pptx`（该目录**已不存在**）
    - `/tmp/cross_domain/cdc_zoonotic_diseases.pptx`（`/tmp` **不持久** ⇒ 从未真跑）
  ⇒ 3 条用例在本机**从未执行过**，其正确性无本地证据（登记见 `.learnings/FEATURE_REQUESTS.md`）。

  实测（v3.24.0 取证）：4 份动医真实课件现位于
  `养禽与禽病防治/M4A1_禽流感/`，实测 `ooxml_lint` 全部 **rc=0**（0 ERROR），
  ⇒ T1 的断言本身成立；**失效的是路径，不是断言**。

设计原则（与本包既有纪律同源）：
  1. **不许把 skip 换个名字**——包外样本是**加分项**，不是**正确性的前提**。
     ⇒ 真实样本可得 ⇒ 用真实样本（证据强度最高）；不可得 ⇒ 用**确定性合成 fixture**
       顶上，断言仍然真跑。**skip 只保留在"合成 fixture 也构造不出来"的极端情形**。
  2. **路径不硬编码单点**——多候选目录 + 环境变量覆盖（`VETMED_SAMPLE_DIR` 等），
     样本搬家/换机不致再次静默 skip。
  3. **不静默降级**——每次降级都**显式登记**降级类型（`real` / `synthetic`），
     由调用方打印，禁"悄悄用弱样本假装测过真样本"。

对外 API：
  resolve_real_samples(...)  -> (list[Path], source_desc)   仅真实样本，不足即返回空
  resolve_samples(...)       -> (list[Path], source_desc)   真实优先，不足补合成
  build_synthetic_deck(...)  -> Path                       确定性合成课件

跨学科样本的**可复现获取方式**（URL / sha256 / 放置路径 / 诚实说明）登记在
  `references/sample-registry.md`
—— 本模块只管「在哪、怎么选」，不管「从哪下载」；两者分开，避免代码里堆事实。
"""
from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import List, Optional, Tuple

# ── 真实样本候选目录（按优先级；每个候选下按文件名 glob 匹配）──
# ⚠ 只读：这些目录属用户工作区，本模块**只读不写**。
#   顺序即优先级：环境变量 > 挂账抽取目录 > 源课件目录。
_WORKSPACE_ROOT = Path(
    "/Users/biodog/Library/CloudStorage/OneDrive-Personal/养禽与禽病防治"
)

#: 候选源：(目录, 文件名 glob 列表, 用途标签)
_SAMPLE_SOURCES: List[Tuple[Path, Tuple[str, ...], str]] = [
    # ① 挂账抽取目录（v3.2.0-batch2 时的原始位置；若被清理则自动落到 ②）
    (
        _WORKSPACE_ROOT / ".harness-runs" / "M4-M7_重诊断" / "sample_extract",
        ("*.pptx",),
        "挂账抽取目录",
    ),
    # ② 源课件目录（**v3.24.0 取证发现的真实位置**）
    #    ⚠ 文件名随用户整理会变，故用多个 glob 兜底（`2.1高致病性…` / `高致病性…` 两种都收）
    (
        _WORKSPACE_ROOT / "M4A1_禽流感",
        (
            "高致病性禽流感防控.pptx",
            "2.1高致病性禽流感防控.pptx",
            "7. 禽流感.pptx",
            "禽流感病毒全链攻防术.pptx",
            "禽流感鉴别诊疗要点.pptx",
        ),
        "源课件目录(M4A1_禽流感)",
    ),
]

#: 环境变量覆盖（换机/自定义样本库时用）
_ENV_SAMPLE_DIR = "VETMED_SAMPLE_DIR"
#: **强制合成**自检开关：置 1 ⇒ 忽略一切真实样本，强制走合成 fixture 分支。
#    用途：验证降级路径本身可跑通（否则降级分支永远没被执行过 = 未验证代码）。
_ENV_FORCE_SYNTH = "VETMED_FORCE_SYNTHETIC"
#: 跨学科样本（v4 通用化用）——**刻意不自动下载**：测试依赖网络会让 CI 不确定。
#    提供方可下载后置于该目录，或设 `VETMED_CROSS_DOMAIN_DIR`。
_ENV_CROSS_DIR = "VETMED_CROSS_DOMAIN_DIR"
_CROSS_DIR_DEFAULT = Path("/tmp/cross_domain")
#: v3.24.0 的**固定文件名**（向后兼容保留）：目录里若恰好有它，仍**优先**选中，
#:    保证既有行为不变（不因本次改造而改变任何已落位样本的选取结果）。
_CROSS_FILENAME = "cdc_zoonotic_diseases.pptx"
#: v3.25.0：目录内**任意** `*.pptx` 的兜底选取——
#:    不同来源的跨学科样本**不必改名**即可落位（原先必须改名成固定文件名）。
#:    ⚠ 跳过 Office 锁文件（`~$` 前缀：0 字节、不是合法 pptx）与隐藏文件。
_CROSS_LOCK_PREFIX = "~$"

#: v3.25.0：跨学科样本的**登记 sha256**（提示级软校验用）。
#:    ⚠ 出处 = `references/sample-registry.md` 的登记条目（唯一真源）。
#:    ⚠ **只提示、不阻断**：样本被替换/损坏属于「环境变化」，不是本包缺陷 ⇒
#:      判红会让换机/换样本的人被无端拦下 ⇒ 违反「假阳性必须为零」。
#:    ⚠ 值**留空则跳过校验**（登记缺失时不应制造假红）。
_CROSS_SHA256_EXPECTED = "ed217fc68e2fc818bf0f1355bcd102e5029bc601ff9bc71f1bdf716f387d0fd9"


def sha256_of(path: Path) -> str:
    """计算文件 sha256（**只读**）。"""
    import hashlib
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def check_cross_domain_sha256(path: Optional[Path]) -> Optional[str]:
    """跨学科样本 sha256 **软校验**（提示级·**不阻断**）。

    返回值：
      · `None` —— 通过 / 无需校验（未登记期望值、路径不存在、强制合成模式）；
      · `str` —— **不一致的说明文本**（调用方负责打印提示，**不得**据此判红）。

    依据：设计文档 v3.25.0 §3.7「改动 4 · sha256 软校验」——
    「不匹配 ⇒ 提示（不阻断）」。
    """
    if force_synthetic() or path is None:
        return None
    expected = _CROSS_SHA256_EXPECTED.strip()
    if not expected:
        return None
    p = Path(path)
    if not p.is_file():
        return None
    try:
        actual = sha256_of(p)
    except OSError as e:
        return f"样本 sha256 计算失败（{p}）：{e}"
    if actual == expected:
        return None
    return (f"样本 sha256 与登记值不一致：\n"
            f"     登记（references/sample-registry.md）: {expected}\n"
            f"     实际（{p.name}）: {actual}\n"
            f"     ⇒ 样本可能被替换或损坏；**本项仅提示、不阻断**（跨学科样本不随包分发，"
            f"换机/换样本属正常情况）")


def force_synthetic() -> bool:
    """是否强制走合成 fixture 分支（自检用）。"""
    return os.environ.get(_ENV_FORCE_SYNTH, "").strip() in ("1", "true", "yes")


def _env_dir(key: str) -> Optional[Path]:
    """读取环境变量指定的目录；未设置或为空 ⇒ None。"""
    raw = os.environ.get(key, "").strip()
    if not raw:
        return None
    return Path(raw).expanduser()


def resolve_real_samples(minimum: int = 1) -> Tuple[List[Path], str]:
    """定位**真实**外部课件样本。

    返回 (样本列表, 来源描述)。样本不足 ``minimum`` 时返回空列表——
    调用方据此决定是否降级到合成 fixture（**不静默**）。

    候选顺序：环境变量目录 → 挂账抽取目录 → 源课件目录。

    ⚠ 强制合成自检（`VETMED_FORCE_SYNTHETIC=1`）⇒ 直接返回空列表。
    """
    if force_synthetic():
        return [], "强制合成模式(VETMED_FORCE_SYNTHETIC=1)"
    found: List[Path] = []
    origins: List[str] = []

    # ① 环境变量优先（用户显式指定，最高优先级）
    env_dir = _env_dir(_ENV_SAMPLE_DIR)
    if env_dir and env_dir.is_dir():
        hits = sorted(env_dir.glob("*.pptx"))
        if hits:
            found.extend(hits)
            origins.append(f"环境变量 {_ENV_SAMPLE_DIR}={env_dir}")

    # ② 挂账抽取目录 → ③ 源课件目录
    for base, globs, label in _SAMPLE_SOURCES:
        if not base.is_dir():
            continue
        hits: List[Path] = []
        for g in globs:
            hits.extend(base.glob(g))
        # 去重（同一文件可能被多个 glob 命中）并保持稳定顺序
        uniq = sorted({p.resolve() for p in hits})
        if uniq:
            found.extend(uniq)
            origins.append(f"{label}({len(uniq)}份)")

    # 全局去重（不同候选目录可能指向同一文件）
    deduped: List[Path] = []
    seen = set()
    for p in found:
        if p in seen:
            continue
        seen.add(p)
        deduped.append(p)

    desc = " + ".join(origins) if origins else "无可用真实样本"
    if len(deduped) < minimum:
        return [], desc
    return deduped, desc


def resolve_cross_domain_sample() -> Tuple[Optional[Path], str]:
    """定位跨学科样本（T11 通用化验证用）。

    返回 (路径, 描述)；不存在 ⇒ (None, 描述)。

    选取规则（v3.25.0，**两段式**）：
      ① **固定文件名优先**（`_CROSS_FILENAME = "cdc_zoonotic_diseases.pptx"`）——
         向后兼容：目录里恰好有它时仍选它，既有行为**完全不变**。
      ② 否则取目录内**任意** `*.pptx`——不同来源的跨学科样本**不必改名**即可落位。
         ⚠ **确定性**：按**文件名**升序排序取第一个（同名再按完整路径排序）。
            排序键写死在此，⇒ 同一目录内容 ⇒ 同一结果（跨机器可复现）。
         ⚠ 跳过 Office 锁文件（`~$` 前缀）与隐藏文件——它们不是合法 pptx。

    ⚠ **不自动下载**：网络依赖会让测试结果不确定；缺失时由调用方降级。
    ⚠ 强制合成自检（`VETMED_FORCE_SYNTHETIC=1`）⇒ 一律视为缺失。
    """
    if force_synthetic():
        return None, "强制合成模式(VETMED_FORCE_SYNTHETIC=1)"
    d = _env_dir(_ENV_CROSS_DIR) or _CROSS_DIR_DEFAULT
    # ① 固定文件名优先（向后兼容）
    p = d / _CROSS_FILENAME
    if p.is_file():
        return p, f"跨学科样本({p})"
    # ② 目录内任意 *.pptx（按文件名排序取第一个 ⇒ 确定性）
    hits = sorted(
        (q for q in d.glob("*.pptx")
         if q.is_file()
         and not q.name.startswith(_CROSS_LOCK_PREFIX)
         and not q.name.startswith(".")),
        key=lambda q: (q.name, str(q)),
    )
    if hits:
        p = hits[0]
        return p, (f"跨学科样本({p}；目录内任意 pptx，按文件名排序取第一个"
                   f"——共 {len(hits)} 个候选)")
    return None, (f"跨学科样本缺失({d}；设 {_ENV_CROSS_DIR} 可覆盖；"
                  f"目录内需至少 1 个 *.pptx)")


# ── 确定性合成 fixture ────────────────────────────────────────
def build_synthetic_deck(path: Path, *, slides: int = 3, with_text: bool = True,
                         with_image: bool = False) -> Path:
    """构造一份**确定性**的最小合法 PPTX（OOXML 合规 ⇒ lint 0 ERROR）。

    用途：外部样本不可得时，让"跑得通"这件事仍被**真实执行**而非 skip。
    参数默认值取最小可用集，保证生成结果稳定可复现。
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    for i in range(max(1, slides)):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        tf = tb.text_frame
        tf.text = f"合成样本 第{i + 1}页" if with_text else f"Synthetic page {i + 1}"
        for para in tf.paragraphs:
            for run in para.runs:
                run.font.size = Pt(14)
        if with_image:
            from PIL import Image  # 仅在需要图片时导入（可选依赖）
            img_path = path.parent / f"_synth_{i}.png"
            Image.new("RGB", (8, 8), (200, 200, 200)).save(img_path)
            slide.shapes.add_picture(str(img_path), Inches(1), Inches(3),
                                     width=Inches(1), height=Inches(1))
            img_path.unlink(missing_ok=True)
    prs.save(str(path))
    return path


def build_synthetic_suite(dest_dir: Optional[Path] = None, count: int = 4) -> List[Path]:
    """构造 ``count`` 份确定性合成课件（模拟"多份真实样本"场景）。"""
    d = Path(dest_dir) if dest_dir else Path(tempfile.mkdtemp(prefix="vetmed_synth_"))
    d.mkdir(parents=True, exist_ok=True)
    decks: List[Path] = []
    for i in range(max(1, count)):
        p = d / f"synth_sample_{i + 1}.pptx"
        build_synthetic_deck(p, slides=2 + i)
        decks.append(p)
    return decks


if __name__ == "__main__":  # pragma: no cover - 人工排查用
    real, desc = resolve_real_samples()
    print(f"真实样本：{len(real)} 份 — {desc}")
    for p in real:
        print(f"  · {p}")
    cd, cdesc = resolve_cross_domain_sample()
    print(f"跨学科样本：{cd or '（缺失）'} — {cdesc}")
