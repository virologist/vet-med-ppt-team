"""inject_notes.py 四前置断言的**行为性**回归护栏（v3.23.0 · 第三轮从零诊断 F3）。

为什么需要这个文件：
  v3.22.0 的 G1 在 `inject_notes.py` 落地了四前置断言
  （字号 font_below_min / 空位与重叠 overlap_exceeded·no_empty_slot / 结果清单），
  那是包内「声明了 8 个月没人读」同型缺陷的收口产物。
  但**实测发现**：`inject_notes.py`（425 行）**从未被任何测试文件触达**——
  唯一提到它的 `test_dead_config_disposition.py` 只在**注释**里引用
  （作为"该键已有代码读取点"的证据），从不 import、从不执行。
  ⇒ 四道防线**没有回归护栏**：任何对配置读取段的重构都可能悄悄废掉它们，
  而**没有一条测试会转红**。这正是历轮「修了但没护栏」的复发温床。

本文件覆盖（每条都有可复现断言）：
  T1 断言② 字号：min_font_pt 高于注入字号 ⇒ 拒写，原因 font_below_min
  T2 断言③ 重叠：注入框压既有文本框 ⇒ 拒写，原因 overlap_exceeded
  T3 断言③ 空位：页内无 ≥注入框尺寸的空白位 ⇒ 拒写，原因 no_empty_slot
  T4 断言④ 结果清单：每页条目含 页码/坐标/字号/重叠值/写或拒+拒因
  T5 正常路径：三断言全过 ⇒ 写入，且备注含【AI草稿·待教师复核】、字号 12pt
  T6 dry-run：判据不降级（仍拒写），且**不写盘**
  T7 断言① 配置读取：四键与 visual-rules.yaml#overlay_injection 一致
  T8 其余零触达脚本的冒烟：--help 可执行、退出码 0
"""
import importlib.util
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
PKG = HERE.parent
INJECT = PKG / "scripts" / "inject_notes.py"


def _load_inject():
    spec = importlib.util.spec_from_file_location("inject_notes_ut", INJECT)
    m = importlib.util.module_from_spec(spec)
    sys.modules["inject_notes_ut"] = m
    spec.loader.exec_module(m)
    return m


def _mk_pptx(path, boxes, slide_w_in=10.0, slide_h_in=7.5):
    """造一份最小课件：boxes = [(x, y, w, h, text), ...]（英寸）。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width = Inches(slide_w_in)
    prs.slide_height = Inches(slide_h_in)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    for (x, y, w, h, text) in boxes:
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.text_frame.text = text
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(18)
    prs.save(str(path))
    return str(path)


class TestInjectNotesAssertions(unittest.TestCase):
    """F3：四前置断言的行为性回归护栏。"""

    @classmethod
    def setUpClass(cls):
        cls.m = _load_inject()
        cls.tmp = Path(tempfile.mkdtemp(prefix="hb-inject-"))

    @classmethod
    def tearDownClass(cls):
        import shutil
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def _deck(self, name, boxes, **kw):
        p = self.tmp / name
        _mk_pptx(p, boxes, **kw)
        return str(p)

    def _bs(self, **over):
        bs = dict(self.m._BS_DEFAULTS)
        bs.update(over)
        return bs

    # ── T1：断言② 字号 ────────────────────────────────────────────────
    def test_T1_font_below_min_rejects_write(self):
        """T1：注入字号（NOTE_FONT_PT=12）< min_font_pt ⇒ 拒写，原因 font_below_min。"""
        deck = self._deck("t1.pptx", [(1, 1, 4, 2, "正文内容")])
        st = self.m.process(deck, cov=1.0, bs=self._bs(min_font_pt=20.0),
                            dry_run=True, inject_box=(6.0, 6.0, 4.0, 1.0))
        self.assertEqual(st["rejected_font"], 1, f"字号断言未触发：{st}")
        entries = st["injection_log"]
        self.assertTrue(any(e["rejected"] == "font_below_min" for e in entries),
                        f"未出现 font_below_min：{entries}")
        self.assertEqual(st["injected"], 0, "被拒写的页不应计入注入数")

    # ── T2：断言③ 重叠 ────────────────────────────────────────────────
    def test_T2_overlap_exceeded_rejects_write(self):
        """T2：注入框与既有文本框重叠比 > max_overlap_ratio ⇒ 拒写 overlap_exceeded。"""
        deck = self._deck("t2.pptx", [(1, 1, 4, 2, "正文内容")])
        # 注入框完全压在既有文本框上 ⇒ 重叠比 = 1.0
        st = self.m.process(deck, cov=1.0, bs=self._bs(max_overlap_ratio=0.15),
                            dry_run=True, inject_box=(1.0, 1.0, 4.0, 1.0))
        self.assertEqual(st["rejected_overlap"], 1, f"重叠断言未触发：{st}")
        entries = st["injection_log"]
        hit = [e for e in entries if e["rejected"] == "overlap_exceeded"]
        self.assertTrue(hit, f"未出现 overlap_exceeded：{entries}")
        self.assertGreater(hit[0]["overlap_ratio"], 0.15,
                            f"重叠比未记录或算错：{hit[0]}")

    # ── T3：断言③ 空位 ────────────────────────────────────────────────
    def test_T3_no_empty_slot_rejects_write(self):
        """T3：页内**容不下**注入框（4×1 英寸）⇒ 拒写，原因 no_empty_slot。

        构造：把页做成 3×1 英寸（**宽度**小于注入框 4 英寸 ⇒ 网格无可用横向落点）。
        ⚠ python-pptx 要求页尺寸 ≥1 英寸，故取 3×1（宽 3 < 注入框宽 4）。
        """
        deck = self._deck("t3.pptx", [(0.2, 0.1, 2.0, 0.6, "短内容")],
                          slide_w_in=3.0, slide_h_in=1.0)
        st = self.m.process(deck, cov=1.0, bs=self._bs(), dry_run=True)
        self.assertEqual(st["rejected_no_empty_slot"], 1, f"空位断言未触发：{st}")
        entries = st["injection_log"]
        self.assertTrue(any(e["rejected"] == "no_empty_slot" for e in entries),
                        f"未出现 no_empty_slot：{entries}")
        self.assertEqual(st["injected"], 0)

    # ── T4：断言④ 结果清单 ────────────────────────────────────────────
    def test_T4_result_manifest_has_all_fields(self):
        """T4：每页清单须含 页码/坐标/字号/重叠值/写或拒+拒因（复核下限用）。"""
        deck = self._deck("t4.pptx", [(1, 1, 4, 2, "正文内容")])
        st = self.m.process(deck, cov=1.0, bs=self._bs(), dry_run=True,
                            inject_box=(6.0, 6.0, 4.0, 1.0))
        entries = st["injection_log"]
        self.assertTrue(entries, "结果清单为空")
        for e in entries:
            for f in ("slide", "font_pt", "inject_box_in", "overlap_ratio",
                      "status", "rejected", "fail_reason"):
                self.assertIn(f, e, f"清单条目缺字段 {f}：{e}")
            self.assertEqual(e["fail_reason"], e["rejected"],
                             "fail_reason 与 rejected 必须同值（两份规格各用其一）")

    # ── T5：正常路径 ──────────────────────────────────────────────────
    def test_T5_normal_path_writes_noted_ai_draft(self):
        """T5：三断言全过 ⇒ 写入备注，且带【AI草稿·待教师复核】标注与 12pt 字号。"""
        from pptx import Presentation
        deck = self._deck("t5.pptx", [(1, 1, 4, 2, "正文内容")])
        st = self.m.process(deck, cov=1.0, bs=self._bs(), dry_run=False,
                            inject_box=(6.0, 6.0, 4.0, 1.0))
        self.assertEqual(st["injected"], 1, f"正常路径未注入：{st}")
        prs = Presentation(deck)
        note = prs.slides[0].notes_slide.notes_text_frame.text
        self.assertIn("AI草稿·待教师复核", note,
                      "备注缺 AI 身份标注（反自评失真纪律）")
        sizes = {r.font.size.pt for p in prs.slides[0].notes_slide.notes_text_frame.paragraphs
                 for r in p.runs if r.font.size is not None}
        self.assertEqual(sizes, {12.0}, f"注入字号不是 12pt：{sizes}")

    # ── T6：dry-run 判据不降级 ────────────────────────────────────────
    def test_T6_dry_run_keeps_criteria_and_writes_nothing(self):
        """T6：--dry-run 下判据**不降级**（仍拒写），且不落盘。"""
        deck = self._deck("t6.pptx", [(1, 1, 4, 2, "正文内容")])
        before = Path(deck).read_bytes()
        st = self.m.process(deck, cov=1.0, bs=self._bs(min_font_pt=20.0),
                            dry_run=True, inject_box=(6.0, 6.0, 4.0, 1.0))
        self.assertEqual(st["rejected_font"], 1, "dry-run 下字号断言被跳过（判据降级）")
        self.assertEqual(Path(deck).read_bytes(), before, "dry-run 下仍写盘")

    # ── T7：断言① 配置读取 ────────────────────────────────────────────
    def test_T7_overlay_injection_keys_match_yaml(self):
        """T7：断言① 读到的四键须与 visual-rules.yaml#overlay_injection 一致。

        （若 yaml 缺键，脚本按铁律「静默回落 = 不设防」告警回落，
          故此处断言"有值且与 yaml 同值或明确告警回落"。）
        """
        import yaml
        cfg = yaml.safe_load(
            (PKG / "config" / "visual-rules.yaml").read_text(encoding="utf-8"))
        want = cfg["overlay_injection"]
        bs = self.m.load_overlay_injection()
        for k in ("min_font_pt", "require_empty_slot", "max_overlap_ratio",
                  "sample_review_ratio"):
            self.assertIn(k, bs, f"断言① 未读到 {k}")
            if want.get(k) is not None:
                self.assertEqual(bs[k], want[k],
                                 f"{k} 读取值与 yaml 不一致：{bs[k]} != {want[k]}")


class TestZeroTouchScriptsSmoke(unittest.TestCase):
    """F3 附带：四个「零测试触达」脚本的冒烟护栏。

    实测（第三轮诊断 P-3）：inject_notes / surgery_deck / fetch_authority /
    check_external_deps 从未被任何测试文件 import 或执行。
    ⇒ 这里给后三者补最小冒烟（--help 可执行、退出码 0），
    至少保证它们不会因语法/导入错误而**静默不可用**。
    """

    def _run_help(self, script):
        return subprocess.run(
            [sys.executable, str(PKG / "scripts" / script), "--help"],
            capture_output=True, text=True, timeout=60)

    def test_smoke_help_exits_zero(self):
        for s in ("surgery_deck.py", "fetch_authority.py", "check_external_deps.py"):
            with self.subTest(script=s):
                r = self._run_help(s)
                self.assertEqual(r.returncode, 0,
                                 f"{s} --help 退出码 {r.returncode}\n{r.stderr[-800:]}")

    def test_fetch_authority_usage_error_not_all_failed(self):
        """F6：未给任何可执行通道输入 ⇒ **用法错误(exit 2)**，不是「全通道失败(exit 3)」。"""
        r = subprocess.run(
            [sys.executable, str(PKG / "scripts" / "fetch_authority.py"),
             "--purpose", "x", "--domain", "animal-medicine"],
            capture_output=True, text=True, timeout=60)
        self.assertNotEqual(r.returncode, 3,
                             "★ 仍在报『全通道失败』——实际一个通道都没跑（假失败）")
        self.assertEqual(r.returncode, 2,
                         f"期望用法错误 exit 2，实际 {r.returncode}")
        self.assertIn("用法错误", r.stderr + r.stdout,
                      "未明示『这是用法错误，不是全通道失败』")


if __name__ == "__main__":
    unittest.main()
